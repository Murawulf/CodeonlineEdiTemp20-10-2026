import re
import pandas as pd
from deep_translator import GoogleTranslator
from openpyxl import Workbook
from openpyxl.utils import get_column_letter
from openpyxl.styles import Font, Alignment


# ===== Step 1: Load and extract keys from .tsx file =====
input_path = "translations.tsx"   # same file you uploaded
with open(input_path, "r", encoding="utf-8") as f:
    text = f.read()

# Extract all string keys from the English section
match = re.search(r'"en"\s*:\s*\{(.*?)\n\s*\}', text, re.S)
if not match:
    raise ValueError("Could not find 'en' section in the file.")

en_section = match.group(1)
pairs = re.findall(r'"(.*?)"\s*:\s*"([^"]*?)"', en_section)
keys = [key.strip() for key, _ in pairs if key.strip()]

print(f"✅ Found {len(keys)} text entries to translate.")


# ===== Step 2: Prepare translation function =====
def translate_text(text, target_lang):
    """Translate a single string safely."""
    try:
        if not text or str(text).strip() == "":
            return ""
        return GoogleTranslator(source="auto", target=target_lang).translate(text)
    except Exception:
        return text


# ===== Step 3: Translate all keys =====
rows = []
for i, text_item in enumerate(keys, 1):
    print(f"🔄 Translating {i}/{len(keys)}: {text_item[:60]}")
    fr = translate_text(text_item, "fr")
    en = translate_text(text_item, "en")
    ar = translate_text(text_item, "ar")
    rows.append({"Texte Source": text_item, "Français": fr, "Anglais": en, "Arabe": ar})

df = pd.DataFrame(rows)


# ===== Step 4: Create Excel file with styling =====
output_path = "translations_auto_translated.xlsx"

wb = Workbook()
ws = wb.active
ws.title = "Traductions"

headers = ["Texte Source", "Français", "Anglais", "Arabe"]
ws.append(headers)

# Bold headers and center them
for col_num, header in enumerate(headers, 1):
    cell = ws.cell(row=1, column=col_num)
    cell.font = Font(bold=True)
    cell.alignment = Alignment(horizontal="center", vertical="center")

# Add data rows
for _, row in df.iterrows():
    ws.append([row["Texte Source"], row["Français"], row["Anglais"], row["Arabe"]])

# Adjust column widths
for col in range(1, len(headers) + 1):
    ws.column_dimensions[get_column_letter(col)].width = 45

wb.save(output_path)
print(f"\n✅ Translation complete! File saved as: {output_path}")
